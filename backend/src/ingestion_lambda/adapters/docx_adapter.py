from pathlib import Path

from docx import Document as DocxDocument

from schemas.models import Chunk


def parse(path: Path, document_id: str) -> list[Chunk]:
    document = DocxDocument(path)
    chunks: list[Chunk] = []
    section: list[str] = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name.lower() if paragraph.style else ""
        is_heading = style.startswith("heading")
        is_list = "list" in style or "bullet" in style
        if is_heading:
            level = style.removeprefix("heading").strip() or "1"
            section = section[: max(0, int(level) - 1)] + [text]
        chunks.append(Chunk(
            id=f"{document_id}:p:{index}", document_id=document_id,
            source_format="docx", text=text,
            element_type="heading" if is_heading else ("list_item" if is_list else "paragraph"),
            section_path=tuple(section),
            location_reference=f"Paragraph: {index + 1}",
        ))
        
    for index, table in enumerate(document.tables):
        md_lines = []
        for i, row in enumerate(table.rows):
            row_data = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            md_lines.append("| " + " | ".join(row_data) + " |")
            if i == 0:
                md_lines.append("|" + "|".join(["---"] * len(row.cells)) + "|")
        text = "\n".join(md_lines)
        if text.strip():
            chunks.append(Chunk(
                id=f"{document_id}:t:{index}", document_id=document_id,
                source_format="docx", text=text,
                element_type="table",
                section_path=(),
                location_reference=f"Table: {index + 1}",
            ))
            
    from .chunk_util import merge_and_split_chunks
    return merge_and_split_chunks(chunks)
